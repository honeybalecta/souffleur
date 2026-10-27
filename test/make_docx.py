"""Kleine Word-Datei für den Import-Test: Überschrift, Absatz, kursiver Absatz."""
import sys
from docx import Document

d = Document()
d.add_heading("Kapitel eins", level=1)
d.add_paragraph("Das ist der erste Satz. Das ist der zweite Satz.")
p = d.add_paragraph()
p.add_run("Hier Bild einblenden.").italic = True
d.add_paragraph("Danach geht es weiter, mit einem kursiven ").add_run("Wort").italic = True
d.save(sys.argv[1])
